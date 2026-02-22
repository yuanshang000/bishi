"""
文档相关API路由
"""

from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form, Query
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from pathlib import Path
import shutil
import hashlib
import logging

from app.schemas.document import DocumentCreate, DocumentUpdate, DocumentResponse
from app.models.document import CourseDocument, DocumentType
from app.models.chapter import Chapter
from app.models.user import User
from app.api.dependencies import get_db, get_current_active_user, require_teacher
from app.services.rag_service import rag_service
from app.services.course_parser import CourseParserService
from app.core.config import settings

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/documents", tags=["文档管理"])


def clean_document_properties(content: str) -> str:
    """
    彻底清理文档内容中的所有元数据和属性信息

    优先使用标记提取法：如果文档中包含[开始标记]和[结束标记]，则只提取标记之间的内容
    否则使用传统清理方法

    Args:
        content: 原始文档内容

    Returns:
        str: 清理后的纯教学内容
    """
    import re

    # 记录原始长度
    original_length = len(content)
    logger.info(f"=== Starting document cleaning, original length: {original_length} chars ===")

    # 🎯 优先方法：使用标记提取
    start_marker = "[开始标记]"
    end_marker = "[结束标记]"

    if start_marker in content and end_marker in content:
        logger.info("Found markers [开始标记] and [结束标记], using marker-based extraction")

        start_idx = content.find(start_marker) + len(start_marker)
        end_idx = content.find(end_marker)

        if start_idx < end_idx:
            extracted_content = content[start_idx:end_idx].strip()

            # 基本清理：只删除多余的空白
            extracted_content = re.sub(r'\n\s*\n\s*\n+', '\n\n', extracted_content)  # 多个空行变成两个
            extracted_content = re.sub(r'[ \t]+', ' ', extracted_content)  # 多个空格变成一个
            extracted_content = re.sub(r'^\s+|\s+$', '', extracted_content, flags=re.MULTILINE)  # 删除每行首尾空格

            logger.info(f"Successfully extracted content between markers: {len(extracted_content)} chars")
            logger.info(f"Content preview: {extracted_content[:200].replace(chr(10), ' ')}...")

            return extracted_content
        else:
            logger.warning("Markers found but in wrong order, falling back to traditional cleaning")
    else:
        logger.info("No markers found, using traditional cleaning method")

    # 🔥 STEP 1: 删除Word XML命名空间和属性标签（这是主要问题）
    # 这些是Word内部的XML标签，不应该被提取出来
    xml_namespace_patterns = [
        r'\(?\s*[owmv]:\s*\w+\s+[^)]*\)?',  # 匹配 (o Author lala)、(w View Web) 等
        r'[owmv]:\w+\s+\w+',  # 匹配 o:Author、w:View 等
        r'o\s+Author\s+\w+',
        r'o\s+LastAuthor\s+\w+',
        r'o\s+Pages\s+\d+',
        r'o\s+Words\s+\d+',
        r'o\s+Characters\s+\d+',
        r'w\s+View\s+\w+',
        r'w\s+BrowserLevel\s+\w+',
        r'm\s+\w+\s+\w+',
        r'v\s+\w+\s+\w+',
    ]

    cleaned_content = content

    for pattern in xml_namespace_patterns:
        before_len = len(cleaned_content)
        cleaned_content = re.sub(pattern, '', cleaned_content, flags=re.IGNORECASE)
        removed = before_len - len(cleaned_content)
        if removed > 0:
            logger.info(f"Removed {removed} chars with pattern: {pattern}")

    # 🔥 STEP 2: 删除HTML/XML标签
    html_patterns = [
        r'<[^>]+>',  # 所有HTML标签
        r'&\w+;',    # HTML实体
    ]

    for pattern in html_patterns:
        cleaned_content = re.sub(pattern, '', cleaned_content, flags=re.IGNORECASE)

    # 🔥 STEP 3: 删除文档属性相关的完整句子或段落
    # 匹配包含这些关键词的整行内容
    property_line_patterns = [
        r'.*?[Aa]uthor.*?\n',
        r'.*?[Ll]ast\s*[Mm]odified.*?\n',
        r'.*?[Pp]ages.*?\n',
        r'.*?[Ww]ords.*?\n',
        r'.*?[Cc]haracters.*?\n',
        r'.*?[Vv]iew\s*[Ww]eb.*?\n',
        r'.*?[Bb]rowser.*?[Ll]evel.*?\n',
        r'.*?[Dd]ocument.*?[Pp]roperties.*?\n',
        r'.*?Microsoft.*?Internet.*?Explorer.*?\n',
        r'.*?single-page document.*?\n',
        r'.*?digital document.*?\n',
        r'.*?originally authored by.*?\n',
        r'.*?configured for.*?web view.*?\n',
        r'.*?compatible with.*?\n',
        r'.*title.*which is just an exclamation mark.*?\n',
    ]

    for pattern in property_line_patterns:
        before_len = len(cleaned_content)
        cleaned_content = re.sub(pattern, '\n', cleaned_content, flags=re.IGNORECASE | re.MULTILINE)
        removed = before_len - len(cleaned_content)
        if removed > 0:
            logger.info(f"Removed property line: {removed} chars")

    # 🔥 STEP 4: 删除传统的文档属性
    traditional_property_patterns = [
        r'作者[:：]\s*\S+',
        r'标题[:：]\s*\S+',
        r'主题[:：]\s*\S+',
        r'公司[:：]\s*\S+',
        r'上次修订者[:：]\s*\S+',
        r'创建时间[:：]\s*[\d\-/: ]+',
        r'修改时间[:：]\s*[\d\-/: ]+',
        r'Microsoft Office Word \d+',
        r'WPS Office',
        r'WPS Writer',
        r'页数[:：]\s*\d+',
        r'字数[:：]\s*\d+',
    ]

    for pattern in traditional_property_patterns:
        cleaned_content = re.sub(pattern, '', cleaned_content, flags=re.IGNORECASE)

    # 🔥 STEP 5: 删除格式相关的描述
    format_patterns = [
        r'Ensure the main headings use.*?font.*?\.',
        r'The document indicates.*?formatted.*?\.',
        r'formatted with.*?font.*?\.',
        r'using.*?point.*?font',
        r'with.*?alignment',
        r'确保.*?使用.*?字体.*?。',
        r'文档显示.*?格式.*?。',
        r'字体.*?粗体',
        r'字体.*?大小',
    ]

    for pattern in format_patterns:
        cleaned_content = re.sub(pattern, '', cleaned_content, flags=re.IGNORECASE | re.DOTALL)

    # 🔥 STEP 6: 清理多余的空白和特殊字符
    cleaned_content = re.sub(r'\n\s*\n\s*\n+', '\n\n', cleaned_content)  # 多个空行变成两个
    cleaned_content = re.sub(r'[ \t]+', ' ', cleaned_content)  # 多个空格变成一个
    cleaned_content = re.sub(r'^\s+|\s+$', '', cleaned_content, flags=re.MULTILINE)  # 删除每行首尾空格
    cleaned_content = cleaned_content.strip()

    # 🔥 STEP 7: 验证清理后的内容
    final_length = len(cleaned_content)
    removed_total = original_length - final_length

    logger.info(f"=== Cleaning completed ===")
    logger.info(f"Original: {original_length} chars")
    logger.info(f"Final: {final_length} chars")
    logger.info(f"Removed: {removed_total} chars ({removed_total*100//original_length if original_length > 0 else 0}%)")

    # 显示清理后的内容预览
    preview_length = min(500, len(cleaned_content))
    preview = cleaned_content[:preview_length].replace('\n', ' ')
    logger.info(f"=== Cleaned content preview (first {preview_length} chars) ===")
    logger.info(f"{preview}")

    return cleaned_content


@router.get("/", response_model=List[DocumentResponse])
def get_documents(
    chapter_id: Optional[int] = None,
    doc_type: Optional[str] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    获取文档列表

    - **chapter_id**: 按章节过滤（可选）
    - **doc_type**: 按类型过滤（可选）
    - **skip**: 跳过数量（分页）
    - **limit**: 返回数量限制
    """
    query = db.query(CourseDocument)

    if chapter_id:
        query = query.filter(CourseDocument.chapter_id == chapter_id)

    if doc_type:
        query = query.filter(CourseDocument.doc_type == doc_type)

    documents = query.offset(skip).limit(limit).all()
    return documents


@router.get("/download/{document_id}")
def download_document(
    document_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    下载文档文件
    """
    logger.info(f"Download request for document_id: {document_id}")

    document = db.query(CourseDocument).filter(CourseDocument.id == document_id).first()
    if not document:
        logger.warning(f"Document not found: {document_id}")
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="文档不存在"
        )

    logger.info(f"Document found: {document.title}, file_path: {document.file_path}")

    # 检查是否有文件路径
    if not document.file_path:
        logger.warning(f"No file_path for document: {document_id}")
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="文档没有关联的文件"
        )

    # 处理路径：替换反斜杠为正斜杠，避免转义字符问题
    normalized_path = document.file_path.replace('\\', '/')
    file_path = Path(normalized_path)
    logger.info(f"Checking file existence: {file_path}")

    # 检查文件是否存在
    if not file_path.exists():
        logger.error(f"File does not exist: {file_path}")
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="文件不存在"
        )

    logger.info(f"Returning file: {document.file_name}")
    # 返回文件
    return FileResponse(
        path=str(file_path),
        filename=document.file_name,
        media_type='application/octet-stream'
    )


@router.get("/{document_id}", response_model=DocumentResponse)
def get_document(
    document_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    获取单个文档详情
    """
    document = db.query(CourseDocument).filter(CourseDocument.id == document_id).first()
    if not document:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="文档不存在"
        )
    return document


@router.post("/", response_model=DocumentResponse, status_code=status.HTTP_201_CREATED)
def create_document(
    document_data: DocumentCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher)
):
    """
    创建文档（仅教师）

    - **chapter_id**: 所属章节ID
    - **title**: 文档标题
    - **doc_type**: 文档类型
    - **content**: 文档内容（可选）
    """
    # 计算内容哈希
    content_hash = None
    if document_data.content:
        content_hash = hashlib.md5(document_data.content.encode()).hexdigest()

    # 创建文档
    document = CourseDocument(
        chapter_id=document_data.chapter_id,
        title=document_data.title,
        doc_type=DocumentType(document_data.doc_type),
        content=document_data.content,
        content_hash=content_hash
    )

    db.add(document)
    db.commit()
    db.refresh(document)

    return document


@router.post("/upload_structure", status_code=status.HTTP_201_CREATED)
async def upload_course_structure(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher)
):
    """
    上传课程结构文档（Markdown），自动生成章节、小节和小目标
    """
    # 检查文件扩展名
    if not file.filename.endswith('.md'):
        raise HTTPException(status_code=400, detail="仅支持 Markdown (.md) 格式的文件用于结构生成。")

    content = await file.read()
    try:
        content_str = content.decode('utf-8')
    except UnicodeDecodeError:
        raise HTTPException(status_code=400, detail="文件必须是 UTF-8 编码。")

    parser = CourseParserService(db)
    try:
        parser.parse_and_create_structure(content_str)
        return {"message": "课程结构生成成功", "filename": file.filename}
    except Exception as e:
        logger.error(f"Failed to parse course structure: {e}")
        raise HTTPException(status_code=500, detail=f"生成课程结构失败: {str(e)}")


@router.get("/template/download", response_class=FileResponse)
def download_template(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    下载课程结构导入模板
    """
    template_path = Path("uploads/template_v2.md")
    if not template_path.exists():
        # 如果文件不存在，创建一个默认的
        content = """# [章节] 第一章：英语基础
## [节] 1.1 基础词汇
### [小目标] 1.1.1 掌握核心名词
**学习内容**：
学习常见的名词，如apple, book, cat等。

**出题要求**：
- 单词拼写
- 图片匹配

**考核指标**：
- 权重：20%
- 及格分：80
- 最少题量：15

### [小目标] 1.1.2 掌握基础动词
**学习内容**：
学习run, eat, sleep等动词。

**考核指标**：
- 权重：30%
- 及格分：60
- 最少题量：10
"""
        template_path.parent.mkdir(parents=True, exist_ok=True)
        with open(template_path, "w", encoding="utf-8") as f:
            f.write(content)
            
    return FileResponse(
        path=str(template_path),
        filename="课程结构导入模板.md",
        media_type='text/markdown'
    )


@router.post("/upload", response_model=DocumentResponse, status_code=status.HTTP_201_CREATED)
async def upload_document(
    file: UploadFile = File(...),
    chapter_id: int = Form(...),
    title: str = Form(...),
    doc_type: str = Form(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher)
):
    """
    上传文档文件（仅教师）

    - **file**: 文档文件（PDF、DOCX、TXT等）
    - **chapter_id**: 所属章节ID
    - **title**: 文档标题
    - **doc_type**: 文档类型
    """
    # 验证章节是否存在
    chapter = db.query(Chapter).filter(Chapter.id == chapter_id).first()
    if not chapter:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="章节不存在"
        )

    # 创建上传目录
    upload_dir = Path(settings.UPLOAD_DIR)
    upload_dir.mkdir(parents=True, exist_ok=True)

    # 生成文件路径
    file_ext = Path(file.filename).suffix
    file_path = upload_dir / f"{chapter_id}_{hashlib.md5(file.filename.encode()).hexdigest()[:8]}{file_ext}"

    # 保存文件
    try:
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"文件上传失败: {str(e)}"
        )

    # 提取文本内容 - 使用多种备用方案
    content = ""
    extraction_method = "unknown"

    try:
        if file_ext == ".txt":
            # 方法1: 直接读取文本文件
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    content = f.read()
                extraction_method = "txt_utf8"
            except UnicodeDecodeError:
                # 尝试其他编码
                for encoding in ['gbk', 'gb2312', 'latin1']:
                    try:
                        with open(file_path, "r", encoding=encoding) as f:
                            content = f.read()
                        extraction_method = f"txt_{encoding}"
                        break
                    except:
                        continue

        elif file_ext == ".pdf":
            # 方法1: 使用pypdf提取PDF
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            content = "\n\n".join([page.extract_text() for page in reader.pages])
            extraction_method = "pypdf"

        elif file_ext in [".docx", ".doc"]:
            # 🎯 简化的Word文档提取策略

            # 方法1: python-docx - 最通用的方法
            if file_ext == ".docx":
                try:
                    from docx import Document

                    logger.info(f"Extracting Word document with python-docx: {file_path}")
                    doc = Document(file_path)

                    # 提取所有段落（不过滤，先看看有什么）
                    all_paragraphs = [para.text for para in doc.paragraphs]
                    logger.info(f"Found {len(all_paragraphs)} paragraphs in document")

                    # 过滤掉空段落
                    paragraphs = [p for p in all_paragraphs if p.strip()]
                    logger.info(f"After filtering empty: {len(paragraphs)} paragraphs")

                    # 如果有段落内容，直接使用
                    if paragraphs:
                        content = "\n\n".join(paragraphs)
                        extraction_method = "python-docx"
                        logger.info(f"✓ Successfully extracted {len(content)} chars from {len(paragraphs)} paragraphs")
                    else:
                        logger.warning("No paragraphs found, trying table extraction...")

                        # 尝试提取表格
                        table_texts = []
                        for table in doc.tables:
                            for row in table.rows:
                                row_text = " | ".join([cell.text.strip() for cell in row.cells])
                                if row_text.strip():
                                    table_texts.append(row_text)

                        if table_texts:
                            content = "\n\n".join(table_texts)
                            extraction_method = "python-docx-tables"
                            logger.info(f"✓ Extracted {len(content)} chars from tables")
                        else:
                            # 段落和表格都为空，可能是特殊格式（如WPS的MHT嵌入）
                            logger.warning("No paragraphs or tables found, checking for embedded content...")
                            raise ValueError("No content in paragraphs or tables")

                except Exception as docx_error:
                    logger.warning(f"python-docx extraction failed: {docx_error}")
                    logger.info("Trying docx2txt as fallback...")
                    content = ""

            # 方法2: docx2txt - 备用方法，通常更宽松
            if not content or len(content.strip()) < 10:
                try:
                    import docx2txt
                    content = docx2txt.process(str(file_path))
                    extraction_method = "docx2txt"

                    if content and len(content.strip()) >= 10:
                        logger.info(f"✓ docx2txt extracted {len(content)} chars")
                    else:
                        raise ValueError("docx2txt returned empty or too short content")

                except Exception as docx2txt_error:
                    logger.warning(f"docx2txt failed: {docx2txt_error}")
                    content = ""

            # 如果前面的方法都失败了
            if not content or len(content.strip()) < 10:
                logger.error(f"❌ All Word extraction methods failed for {file_path}")
                logger.error("Possible reasons:")
                logger.error("  1. Document is corrupted")
                logger.error("  2. Document uses unsupported format (e.g., WPS Office specific format)")
                logger.error("  3. Document contains only images or embedded objects")
                logger.error("Suggestion: Convert to PDF or copy text to plain .txt file")

    except Exception as e:
        # 如果所有方法都失败
        logger.error(f"All extraction methods failed for {file_path}: {e}")

    # 如果最终还是没有提取到有效内容
    if not content or len(content.strip()) < 10:
        # 删除已上传的文件
        try:
            if file_path.exists():
                file_path.unlink()
        except Exception as del_error:
            logger.error(f"Failed to delete file after extraction error: {del_error}")

        # 返回详细错误信息
        error_msg = f"文件上传成功，但无法提取有效内容。已尝试多种解析方法。"
        if file_ext in [".docx", ".doc"]:
            error_msg += " 请尝试：1. 用Word重新保存文件  2. 转换为PDF格式  3. 复制内容到新文档"
        elif file_ext == ".pdf":
            error_msg += " 请确保PDF文件未加密且未损坏。"
        else:
            error_msg += f" 不支持的文件格式：{file_ext}"

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=error_msg
        )

    # 成功提取内容
    logger.info(f"Content extracted successfully: method={extraction_method}, length={len(content)} chars")

    # 详细记录提取的内容预览（前500字符）
    content_preview = content[:500].replace('\n', ' ')
    logger.info(f"=== Extracted Content Preview (first 500 chars) ===")
    logger.info(f"{content_preview}")

    # ⚠️ 重要：上传时保存原始内容，不进行清理
    # 清理操作将在向量化时进行，这样可以：
    # 1. 保留文档中的标记（[开始标记]和[结束标记]）
    # 2. 允许随时调整清理逻辑而不需要重新上传
    # 3. 预览功能可以正确显示原始内容和清理后的差异

    # 检查原始内容是否有效
    if not content or len(content.strip()) < 50:
        logger.error("Content too short or empty!")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="文档内容过短或为空。文件可能损坏或没有可提取的文本内容。"
        )

    # 计算内容哈希（基于原始内容）
    content_hash = hashlib.md5(content.encode()).hexdigest()

    # 创建文档记录（使用正斜杠路径）
    document = CourseDocument(
        chapter_id=chapter_id,
        title=title,
        doc_type=DocumentType(doc_type),
        file_path=str(file_path).replace('\\', '/'),
        file_name=file.filename,
        file_size=file_path.stat().st_size,
        content=content,
        content_hash=content_hash
    )

    db.add(document)
    db.commit()
    db.refresh(document)

    logger.info(f"Document uploaded successfully: ID={document.id}, extraction_method={extraction_method}, content_length={len(content)}")
    return document


@router.post("/{document_id}/re-extract", status_code=status.HTTP_200_OK)
def re_extract_document_content(
    document_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher)
):
    """
    重新提取文档内容（仅教师）

    用于修复上传时内容提取失败的文档
    """
    document = db.query(CourseDocument).filter(CourseDocument.id == document_id).first()
    if not document:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="文档不存在"
        )

    if not document.file_path:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="该文档没有关联的文件"
        )

    # 获取文件路径
    file_path = Path(document.file_path.replace('\\', '/'))
    if not file_path.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="文档文件不存在"
        )

    # 使用相同的提取逻辑
    file_ext = file_path.suffix.lower()
    content = ""
    extraction_method = "unknown"

    # 复用上传时的提取逻辑
    try:
        if file_ext == ".txt":
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    content = f.read()
                extraction_method = "txt_utf8"
            except UnicodeDecodeError:
                for encoding in ['gbk', 'gb2312', 'latin1']:
                    try:
                        with open(file_path, "r", encoding=encoding) as f:
                            content = f.read()
                        extraction_method = f"txt_{encoding}"
                        break
                    except:
                        continue

        elif file_ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            content = "\n\n".join([page.extract_text() for page in reader.pages])
            extraction_method = "pypdf"

        elif file_ext in [".docx", ".doc"]:
            # 🎯 简化的Word重新提取

            # 方法1: python-docx
            if file_ext == ".docx":
                try:
                    from docx import Document

                    logger.info(f"Re-extracting with python-docx: {file_path}")
                    doc = Document(file_path)

                    # 提取段落
                    all_paragraphs = [para.text for para in doc.paragraphs]
                    paragraphs = [p for p in all_paragraphs if p.strip()]
                    logger.info(f"Re-extracted {len(paragraphs)} paragraphs")

                    if paragraphs:
                        content = "\n\n".join(paragraphs)
                        extraction_method = "python-docx"
                        logger.info(f"✓ Re-extracted {len(content)} chars")
                    else:
                        # 尝试表格
                        table_texts = []
                        for table in doc.tables:
                            for row in table.rows:
                                row_text = " | ".join([cell.text.strip() for cell in row.cells])
                                if row_text.strip():
                                    table_texts.append(row_text)

                        if table_texts:
                            content = "\n\n".join(table_texts)
                            extraction_method = "python-docx-tables"
                        else:
                            raise ValueError("No content found")

                except Exception as e:
                    logger.warning(f"python-docx re-extraction failed: {e}")
                    content = ""

            # 方法2: docx2txt
            if not content or len(content.strip()) < 10:
                try:
                    import docx2txt
                    content = docx2txt.process(str(file_path))
                    extraction_method = "docx2txt"
                    logger.info(f"✓ docx2txt re-extracted {len(content)} chars")
                except Exception as e:
                    logger.warning(f"docx2txt re-extraction failed: {e}")
                    content = ""

            # 如果都失败
            if not content or len(content.strip()) < 10:
                logger.error(f"❌ All Word re-extraction methods failed")
                logger.error("Suggestion: Convert to PDF or plain text format")

    except Exception as e:
        logger.error(f"Re-extraction failed: {e}")

    if not content or len(content.strip()) < 10:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="重新提取失败：无法从文件中提取有效内容。建议删除后上传新文件。"
        )

    # 详细记录重新提取的内容预览
    content_preview = content[:500].replace('\n', ' ')
    logger.info(f"=== Re-extracted Content Preview (first 500 chars) ===")
    logger.info(f"{content_preview}")

    # ⚠️ 重要：保存原始内容，不进行清理
    # 清理操作将在向量化时进行

    # 检查原始内容是否有效
    if not content or len(content.strip()) < 50:
        logger.error("Re-extracted content too short or empty!")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="重新提取的内容过短或为空。文件可能损坏或没有可提取的文本内容。"
        )

    # 更新文档内容（保存原始内容）
    document.content = content
    document.content_hash = hashlib.md5(content.encode()).hexdigest()
    db.commit()

    logger.info(f"Document {document_id} content re-extracted: method={extraction_method}, length={len(content)}")
    return {
        "message": "内容重新提取成功",
        "document_id": document_id,
        "extraction_method": extraction_method,
        "content_length": len(content)
    }


@router.post("/{document_id}/vectorize", status_code=status.HTTP_200_OK)
def vectorize_document(
    document_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher)
):
    """
    向量化文档（仅教师）
    """
    # 先检查文档是否存在以及是否有内容
    document = db.query(CourseDocument).filter(CourseDocument.id == document_id).first()
    if not document:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="文档不存在"
        )

    if not document.content or len(document.content.strip()) < 10:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="该文档内容为空或过短，无法向量化。请先使用'重新提取内容'功能，或删除后重新上传。"
        )

    # 如果已经向量化，先删除旧向量
    if document.is_vectorized:
        logger.info(f"Document {document_id} is already vectorized, deleting old vectors first")
        rag_service.delete_document_vectors(document_id)
        # 重置向量化状态
        document.is_vectorized = 0
        document.vector_count = 0
        document.chroma_ids = None
        db.commit()

    success = rag_service.vectorize_document(db, document_id)

    if not success:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="向量化失败，请检查文档内容是否正确，或稍后重试"
        )

    return {"message": "向量化成功", "document_id": document_id}


@router.get("/{document_id}/vectorized-content-preview", status_code=status.HTTP_200_OK)
def get_vectorized_content_preview(
    document_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    获取文档向量化内容的预览（所有用户可见）

    返回：
    - raw_content_preview: 原始文档内容预览（前1000字符）
    - cleaned_content_preview: 清理后的内容预览（前1000字符）
    - cleaned_content_length: 清理后的内容总长度
    - is_marker_based: 是否使用了标记提取
    - vector_count: 向量数量
    - chunks_preview: 每个向量块的预览（前200字符）
    """
    document = db.query(CourseDocument).filter(CourseDocument.id == document_id).first()
    if not document:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="文档不存在"
        )

    if not document.content:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="该文档没有内容"
        )

    # 原始内容预览
    raw_content = document.content
    raw_preview = raw_content[:1000]

    # 清理后的内容
    cleaned_content = clean_document_properties(raw_content)
    cleaned_preview = cleaned_content[:1000]

    # 检查是否使用了标记提取
    has_start_marker = "[开始标记]" in raw_content
    has_end_marker = "[结束标记]" in raw_content
    is_marker_based = has_start_marker and has_end_marker

    # 调试信息：记录标记检测结果
    logger.info(f"Marker detection for document {document_id}:")
    logger.info(f"  - Has [开始标记]: {has_start_marker}")
    logger.info(f"  - Has [结束标记]: {has_end_marker}")
    logger.info(f"  - Is marker based: {is_marker_based}")
    logger.info(f"  - Raw content first 500 chars: {raw_content[:500]}")

    # 获取向量块预览
    chunks_preview = []
    if document.is_vectorized:
        # 重新分块以显示实际向量化的块
        from app.services.rag_service import rag_service
        chunks = rag_service.split_text(cleaned_content)
        chunks_preview = [
            {
                "index": i,
                "length": len(chunk),
                "preview": chunk[:200] + ("..." if len(chunk) > 200 else "")
            }
            for i, chunk in enumerate(chunks[:10])  # 最多显示前10个块
        ]

    return {
        "document_id": document_id,
        "title": document.title,
        "raw_content_length": len(raw_content),
        "raw_content_preview": raw_preview,
        "cleaned_content_length": len(cleaned_content),
        "cleaned_content_preview": cleaned_preview,
        "is_marker_based": is_marker_based,
        "has_start_marker": has_start_marker,
        "has_end_marker": has_end_marker,
        "is_vectorized": bool(document.is_vectorized),
        "vector_count": document.vector_count or 0,
        "chunks_preview": chunks_preview
    }


@router.get("/debug/vectors-info", status_code=status.HTTP_200_OK)
def get_vectors_info(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher)
):
    """
    调试接口：查看向量库中的所有向量信息（仅教师）
    """
    try:
        # 获取 ChromaDB 集合的所有数据
        all_data = rag_service.collection.get(
            include=["metadatas"]
        )

        # 统计信息
        total_vectors = len(all_data['ids']) if all_data['ids'] else 0

        # 按文档ID分组
        doc_vectors = {}
        if all_data['metadatas']:
            for metadata in all_data['metadatas']:
                doc_id = metadata.get('document_id')
                if doc_id not in doc_vectors:
                    doc_vectors[doc_id] = {
                        'document_id': doc_id,
                        'title': metadata.get('title'),
                        'doc_type': metadata.get('doc_type'),
                        'chapter_id': metadata.get('chapter_id'),
                        'vector_count': 0
                    }
                doc_vectors[doc_id]['vector_count'] += 1

        # 查询数据库中的文档信息
        db_documents = db.query(CourseDocument).all()
        db_doc_info = {
            doc.id: {
                'id': doc.id,
                'title': doc.title,
                'chapter_id': doc.chapter_id,
                'is_vectorized': doc.is_vectorized,
                'vector_count': doc.vector_count,
                'exists_in_db': True
            } for doc in db_documents
        }

        # 合并信息，找出不一致的地方
        inconsistencies = []
        for doc_id, vec_info in doc_vectors.items():
            db_info = db_doc_info.get(doc_id)
            if not db_info:
                inconsistencies.append({
                    'document_id': doc_id,
                    'issue': 'Vector exists but document not in database',
                    'vector_info': vec_info
                })
            elif not db_info['is_vectorized']:
                inconsistencies.append({
                    'document_id': doc_id,
                    'issue': 'Vector exists but document.is_vectorized=False',
                    'vector_count_in_chroma': vec_info['vector_count'],
                    'vector_count_in_db': db_info['vector_count']
                })

        # 找出数据库中标记为已向量化但向量库中没有的文档
        for doc_id, db_info in db_doc_info.items():
            if db_info['is_vectorized'] and doc_id not in doc_vectors:
                inconsistencies.append({
                    'document_id': doc_id,
                    'issue': 'Document marked as vectorized but no vectors in ChromaDB',
                    'db_info': db_info
                })

        return {
            'total_vectors': total_vectors,
            'documents_with_vectors': len(doc_vectors),
            'vectors_by_document': list(doc_vectors.values()),
            'database_documents': list(db_doc_info.values()),
            'inconsistencies': inconsistencies
        }

    except Exception as e:
        logger.error(f"Error getting vectors info: {e}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"获取向量信息失败: {str(e)}"
        )


@router.post("/rebuild-index", status_code=status.HTTP_200_OK)
def rebuild_index(
    chapter_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher)
):
    """
    重建向量索引（仅教师）

    清除指定章节（或所有章节）的向量数据，然后重新向量化所有文档
    """
    try:
        # 第一步：先删除所有孤儿向量（向量库中存在但数据库中不存在的）
        logger.info("Step 1: Deleting orphan vectors...")
        orphan_deleted = rag_service.delete_orphan_vectors(db)
        logger.info(f"Deleted {orphan_deleted} orphan document vectors")

        # 获取需要处理的文档
        query = db.query(CourseDocument)
        if chapter_id:
            query = query.filter(CourseDocument.chapter_id == chapter_id)
        documents = query.all()

        if not documents:
            return {
                "message": "没有找到需要处理的文档",
                "chapter_id": chapter_id,
                "orphan_deleted": orphan_deleted
            }

        # 第二步：删除所有相关文档的向量
        logger.info("Step 2: Deleting vectors for existing documents...")
        deleted_count = 0
        for doc in documents:
            if doc.is_vectorized:
                success = rag_service.delete_document_vectors(doc.id)
                if success:
                    deleted_count += 1
                # 重置数据库中的向量化状态
                doc.is_vectorized = 0
                doc.vector_count = 0
                doc.chroma_ids = None

        db.commit()
        logger.info(f"Deleted vectors for {deleted_count} existing documents")

        # 第三步：重新向量化所有有内容的文档
        logger.info("Step 3: Re-vectorizing documents...")
        vectorized_count = 0
        failed_docs = []
        for doc in documents:
            if doc.content and len(doc.content.strip()) >= 10:
                success = rag_service.vectorize_document(db, doc.id)
                if success:
                    vectorized_count += 1
                    logger.info(f"Successfully vectorized document {doc.id}: {doc.title}")
                else:
                    failed_docs.append({'id': doc.id, 'title': doc.title})
                    logger.error(f"Failed to vectorize document {doc.id}: {doc.title}")

        return {
            "message": "索引重建完成",
            "chapter_id": chapter_id,
            "total_documents": len(documents),
            "orphan_deleted": orphan_deleted,
            "deleted_vectors": deleted_count,
            "vectorized_documents": vectorized_count,
            "failed_documents": failed_docs
        }

    except Exception as e:
        db.rollback()
        logger.error(f"Error rebuilding index: {e}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"重建索引失败: {str(e)}"
        )


@router.post("/clear-orphan-vectors", status_code=status.HTTP_200_OK)
def clear_orphan_vectors(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher)
):
    """
    清理孤儿向量（仅教师）

    删除向量库中存在但数据库中已删除的文档的向量
    """
    try:
        deleted_count = rag_service.delete_orphan_vectors(db)
        return {
            "message": f"成功清理 {deleted_count} 个孤儿文档的向量",
            "deleted_count": deleted_count
        }
    except Exception as e:
        logger.error(f"Error clearing orphan vectors: {e}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"清理孤儿向量失败: {str(e)}"
        )


@router.post("/batch/re-vectorize", status_code=status.HTTP_200_OK)
def batch_re_vectorize_documents(
    chapter_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher)
):
    """
    批量重新向量化文档（仅教师）

    可选择指定章节ID，只重新向量化该章节的文档。
    如果不指定章节ID，则重新向量化所有已向量化的文档。

    这个功能用于：
    1. 修复元数据类型问题
    2. 更新向量化算法后重新生成向量
    3. 修复章节过滤问题
    """
    try:
        # 查询需要重新向量化的文档
        query = db.query(CourseDocument).filter(CourseDocument.is_vectorized == 1)
        if chapter_id is not None:
            query = query.filter(CourseDocument.chapter_id == chapter_id)

        documents = query.all()

        if not documents:
            return {
                "message": "没有需要重新向量化的文档",
                "total": 0,
                "success": 0,
                "failed": 0
            }

        success_count = 0
        failed_count = 0
        failed_docs = []

        for doc in documents:
            try:
                logger.info(f"Re-vectorizing document {doc.id}: {doc.title} (Chapter {doc.chapter_id})")

                # 删除旧向量
                rag_service.delete_document_vectors(doc.id)

                # 重置向量化状态
                doc.is_vectorized = 0
                doc.vector_count = 0
                doc.chroma_ids = None
                db.commit()

                # 重新向量化
                success = rag_service.vectorize_document(db, doc.id)

                if success:
                    success_count += 1
                    logger.info(f"Successfully re-vectorized document {doc.id}")
                else:
                    failed_count += 1
                    failed_docs.append({"id": doc.id, "title": doc.title, "error": "向量化失败"})
                    logger.error(f"Failed to re-vectorize document {doc.id}")

            except Exception as e:
                failed_count += 1
                failed_docs.append({"id": doc.id, "title": doc.title, "error": str(e)})
                logger.error(f"Error re-vectorizing document {doc.id}: {e}")

        message = f"重新向量化完成: 成功 {success_count}/{len(documents)}"
        if chapter_id:
            message += f" (章节 {chapter_id})"

        return {
            "message": message,
            "total": len(documents),
            "success": success_count,
            "failed": failed_count,
            "failed_documents": failed_docs
        }

    except Exception as e:
        logger.error(f"Error in batch re-vectorization: {e}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"批量重新向量化失败: {str(e)}"
        )


@router.delete("/vectors/clear-all", status_code=status.HTTP_200_OK)
def clear_all_vectors(
    confirm: str = Query(..., description="输入'DELETE_ALL_VECTORS'确认删除"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher)
):
    """
    清空向量库中的所有向量数据（仅教师）

    ⚠️ 警告：这是一个危险操作！会删除所有向量数据！

    使用场景：
    1. 向量数据损坏需要重建
    2. 章节过滤失效，需要清空重建
    3. 测试环境重置

    需要输入确认字符串: DELETE_ALL_VECTORS
    """
    # 安全确认
    if confirm != "DELETE_ALL_VECTORS":
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="请输入正确的确认字符串: DELETE_ALL_VECTORS"
        )

    try:
        # 清空所有向量
        deleted_count = rag_service.clear_all_vectors()

        # 重置所有文档的向量化状态
        db.query(CourseDocument).update({
            "is_vectorized": 0,
            "vector_count": 0,
            "chroma_ids": None
        })
        db.commit()

        logger.warning(f"All vectors cleared by user {current_user.username}: {deleted_count} vectors deleted")

        return {
            "message": f"成功清空 {deleted_count} 个向量，并重置所有文档的向量化状态",
            "deleted_count": deleted_count,
            "warning": "请重新向量化所需的文档"
        }

    except Exception as e:
        logger.error(f"Error clearing all vectors: {e}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"清空向量失败: {str(e)}"
        )



@router.delete("/{document_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_document(
    document_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_teacher)
):
    """
    删除文档（仅教师）
    """
    document = db.query(CourseDocument).filter(CourseDocument.id == document_id).first()
    if not document:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="文档不存在"
        )

    logger.info(f"Deleting document: ID={document_id}, title='{document.title}', is_vectorized={document.is_vectorized}")

    # 删除向量
    if document.is_vectorized:
        logger.info(f"Deleting vectors for document {document_id}")
        success = rag_service.delete_document_vectors(document_id)
        if success:
            logger.info(f"Successfully deleted vectors for document {document_id}")
        else:
            logger.warning(f"Failed to delete vectors for document {document_id}")

    # 删除文件
    if document.file_path:
        # 处理路径：替换反斜杠为正斜杠
        normalized_path = document.file_path.replace('\\', '/')
        file_path = Path(normalized_path)
        if file_path.exists():
            file_path.unlink()
            logger.info(f"Deleted file: {file_path}")

    # 删除数据库记录
    db.delete(document)
    db.commit()

    logger.info(f"Document {document_id} deleted successfully from database")
    return None
